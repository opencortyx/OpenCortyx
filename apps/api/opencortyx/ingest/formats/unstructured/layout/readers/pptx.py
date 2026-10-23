"""PowerPoint (.pptx/.pptm) with python-pptx, slide by slide (the slide number is the page): the slide
title is a title, so the slide's content nests under it; bulleted or indented paragraphs are list
items, the rest text; tables are grids, pictures pictures (by their alt text)."""
from __future__ import annotations

import io
from typing import Optional

from ..assembler import LayoutAssembler
from ..engine import grid_text


class PptxReader:
    name = "pptx"
    extensions = ("pptx", "pptm", "potx")

    def read(self, content: bytes, filename: str, asm: LayoutAssembler,
             max_pages: Optional[int] = None) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
        prs = Presentation(io.BytesIO(content))
        from .native import office_core
        asm.metadata.update(office_core(prs.core_properties))
        for page_no, slide in enumerate(prs.slides, start=1):
            if max_pages and page_no > max_pages:
                break
            shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
            title = None
            for s in shapes:
                if s.is_placeholder and s.placeholder_format.type in (PP_PLACEHOLDER.TITLE,
                                                                       PP_PLACEHOLDER.CENTER_TITLE):
                    title = s
                    break
            if title is not None and title.has_text_frame and title.text_frame.text.strip():
                asm.add_element(raw_label="title", text=title.text_frame.text.strip(), page_no=page_no, level=0)
            for s in shapes:
                if s is title:
                    continue
                if s.has_text_frame:
                    for para in s.text_frame.paragraphs:
                        text = (para.text or "").strip()
                        if not text:
                            continue
                        bulleted = para.level > 0 or para._p.find(".//{*}buChar") is not None \
                            or para._p.find(".//{*}buAutoNum") is not None
                        asm.add_element(raw_label="list_item" if bulleted else "text", text=text, page_no=page_no)
                elif getattr(s, "has_table", False) and s.has_table:
                    grid = [[c.text.strip() for c in row.cells] for row in s.table.rows]
                    grid = [r for r in grid if any(r)]
                    if grid:
                        asm.add_element(raw_label="table", text=grid_text(grid), page_no=page_no, grid=grid)
                elif s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    descr = (s._element.xpath(".//p:cNvPr/@descr") or [""])[0] or s.name
                    asm.add_element(raw_label="picture", text=descr, page_no=page_no)

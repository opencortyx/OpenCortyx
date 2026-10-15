"""Test data: real files in every catalogued format, generated from one value per variant.

A variant is defined ONCE as plain data (rows, a JSON value, a document body) and serialised into each
format with that format's own writer, so a test can check that what comes out of the reader is what
went in. Formats with no writer available are reported BLOCKED by the tests, never faked.
"""
from __future__ import annotations

import io
import json
import zipfile

# ── semi-structured values ───────────────────────────────────────────────────────────────────────


def _deep(n):
    v = {"leaf": "bottom"}
    for i in range(n):
        v = {f"level{i}": v}
    return v


#: name -> (value, expectation). `accept` = must ingest; `reject` = must fail with a ParseError.
SEMI_VARIANTS = {
    "flat_object": ({"id": "A-1", "name": "Alpha", "qty": 3, "active": True}, "accept"),
    "nested_object": ({"order": {"order_id": "O-1", "customer": {"name": "Ann", "address": {"city": "Oslo"}}}}, "accept"),
    "record_array": ([{"id": f"R-{i}", "name": f"row {i}", "n": i} for i in range(5)], "accept"),
    "scalar_list": ({"doc": {"tags": ["alpha", "beta", "gamma"]}}, "accept"),
    "mixed_list": ({"items": [{"k": "v1"}, {"k": "v2", "extra": {"z": "deep"}}]}, "accept"),
    "deep_nesting_40": (_deep(40), "accept"),
    "wide_300_keys": ({"row": {f"field_{i}": f"value_{i}" for i in range(300)}}, "accept"),
    "unicode": ({"doc": {"名前": "東京", "emoji": "\U0001F600 ok", "rtl": "مرحبا", "accent": "Café"}}, "accept"),
    "null_values": ({"doc": {"a": None, "b": "kept", "c": {"d": None, "e": "x"}}}, "accept"),
    "numbers_boundary": ({"n": {"big": 9007199254740993, "neg": -1, "tiny": 1.5e-10, "zero": 0}}, "accept"),
    "long_string_100k": ({"doc": {"text": "word " * 20000}}, "accept"),
    "large_array_3000": ([{"id": f"L-{i}", "v": i} for i in range(3000)], "accept"),
    "duplicate_ids": ([{"id": "SAME", "v": 1}, {"id": "SAME", "v": 2}], "accept"),
    "empty_object": ({}, "reject"),
    "empty_array": ([], "reject"),
}
#: raw-bytes variants that are invalid in every format
SEMI_RAW = {
    "empty_bytes": b"",
    "whitespace_only": b"   \n\t  \n",
    "binary_garbage": bytes(range(256)) * 4,
}

SEMI_FORMATS = ("json", "geojson", "har", "fhir", "jsonl", "ndjson", "xml", "rss", "atom", "yaml", "yml", "toml")


def _xml_name(k: str) -> str:
    k = "".join(c if (c.isalnum() or c in "_-.") else "_" for c in str(k))
    return k if k and (k[0].isalpha() or k[0] == "_") else f"_{k}"


def _to_xml(value, tag: str = "root") -> str:
    from xml.sax.saxutils import escape
    t = _xml_name(tag)
    if isinstance(value, dict):
        return f"<{t}>" + "".join(_to_xml(v, k) for k, v in value.items()) + f"</{t}>"
    if isinstance(value, list):
        return f"<{t}>" + "".join(_to_xml(v, "item") for v in value) + f"</{t}>"
    if value is None:
        return f"<{t}/>"
    return f"<{t}>{escape(str(value).lower() if isinstance(value, bool) else str(value))}</{t}>"


def _toml_value(v) -> str:
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(v)
    return json.dumps(str(v), ensure_ascii=False)


def _to_toml(value: dict) -> str:
    """Minimal TOML writer: scalars, nested tables, arrays of tables, arrays of scalars."""
    lines: list[str] = []

    def table(d: dict, path: list[str]):
        scal = {k: v for k, v in d.items() if not isinstance(v, (dict, list)) or
                (isinstance(v, list) and all(not isinstance(x, (dict, list)) for x in v))}
        for k, v in scal.items():
            key = json.dumps(str(k), ensure_ascii=False)
            lines.append(f"{key} = [{', '.join(_toml_value(x) for x in v)}]" if isinstance(v, list)
                         else f"{key} = {_toml_value(v)}")
        for k, v in d.items():
            if k in scal:
                continue
            p = path + [json.dumps(str(k), ensure_ascii=False)]
            if isinstance(v, dict):
                lines.append(f"\n[{'.'.join(p)}]")
                table(v, p)
            else:
                for el in v:
                    lines.append(f"\n[[{'.'.join(p)}]]")
                    table(el if isinstance(el, dict) else {"value": el}, p)
    table(value, [])
    return "\n".join(lines) + "\n"


def toml_applicable(value) -> bool:
    """TOML has no null and no top-level array; those variants do not exist in TOML."""
    def has_null(v):
        if v is None:
            return True
        if isinstance(v, dict):
            return any(has_null(x) for x in v.values())
        if isinstance(v, list):
            return any(has_null(x) for x in v)
        return False
    return isinstance(value, dict) and not has_null(value) and not _big_int(value)


def _big_int(v) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, int):
        return abs(v) >= 2 ** 63
    if isinstance(v, dict):
        return any(_big_int(x) for x in v.values())
    if isinstance(v, list):
        return any(_big_int(x) for x in v)
    return False


def semi_bytes(fmt: str, value) -> bytes | None:
    """The value written in `fmt`, or None when the format cannot express it."""
    import yaml
    if fmt in ("json", "geojson", "har", "fhir"):
        return json.dumps(value, ensure_ascii=False).encode()
    if fmt in ("jsonl", "ndjson"):
        rows = value if isinstance(value, list) else [value]
        return "\n".join(json.dumps(r, ensure_ascii=False) for r in rows).encode() + (b"\n" if rows else b"")
    if fmt in ("xml", "rss", "atom"):
        return ('<?xml version="1.0" encoding="UTF-8"?>' + _to_xml(value, "records" if isinstance(value, list) else "root")).encode()
    if fmt in ("yaml", "yml"):
        return yaml.safe_dump(value, allow_unicode=True, sort_keys=False).encode()
    if fmt == "toml":
        return _to_toml(value).encode() if toml_applicable(value) else None
    raise ValueError(fmt)


def truncated(b: bytes) -> bytes:
    """A syntactically broken copy: cut mid-structure."""
    return b[: max(1, len(b) * 2 // 3)]


def scalar_leaves(value, depth: int = 0, max_depth: int = 10**6):
    """Every scalar in the value within `max_depth` levels, as strings (what fidelity checks look for)."""
    if depth > max_depth:
        return
    if isinstance(value, dict):
        for v in value.values():
            yield from scalar_leaves(v, depth + 1, max_depth)
    elif isinstance(value, list):
        for v in value:
            yield from scalar_leaves(v, depth + 1, max_depth)
    elif value is not None:
        yield str(value).lower() if isinstance(value, bool) else str(value)


# ── tables ───────────────────────────────────────────────────────────────────────────────────────


def _rows(n, cols=("id", "name", "amount")):
    return [{c: (i if c == "id" else f"{c}_{i}" if c == "name" else round(i * 1.5, 2)) for c in cols}
            for i in range(n)]


#: name -> (columns, rows, expectation)
TABLE_VARIANTS = {
    "basic_5_rows": (["id", "name", "amount"], _rows(5), "accept"),
    "one_row": (["id", "name", "amount"], _rows(1), "accept"),
    "rows_1000": (["id", "name", "amount"], _rows(1000), "accept"),
    "rows_1001": (["id", "name", "amount"], _rows(1001), "accept"),
    "rows_50000": (["id", "name", "amount"], _rows(50000), "accept"),
    "wide_120_cols": ([f"c{i}" for i in range(120)], [{f"c{i}": i * r for i in range(120)} for r in range(3)], "accept"),
    "nulls": (["id", "name", "amount"], [{"id": 1, "name": None, "amount": 2.0}, {"id": 2, "name": "b", "amount": None}], "accept"),
    "unicode_values": (["id", "city"], [{"id": 1, "city": "東京"}, {"id": 2, "city": "Zürich"}, {"id": 3, "city": "São Paulo"}], "accept"),
    "unicode_headers": (["名前", "città"], [{"名前": "a", "città": "b"}, {"名前": "c", "città": "d"}], "accept"),
    "text_with_delimiters": (["id", "note"], [{"id": 1, "note": "a, b; c | d"}, {"id": 2, "note": 'quote "x"'}], "accept"),
    "booleans_dates": (["id", "ok", "day"], [{"id": 1, "ok": True, "day": "2024-01-31"}, {"id": 2, "ok": False, "day": "2025-12-01"}], "accept"),
    "header_only": (["id", "name"], [], "accept_empty"),
    "empty_file": ([], [], "reject"),
}
TABLE_FORMATS = ("csv", "tsv", "parquet", "orc", "feather", "arrow", "avro", "xlsx", "xlsm", "xls")


def _delimited(cols, rows, sep):
    import csv
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=sep, lineterminator="\n")
    if cols:
        w.writerow(cols)
    for r in rows:
        w.writerow(["" if r.get(c) is None else r.get(c) for c in cols])
    return buf.getvalue().encode()


def _arrow_table(cols, rows):
    import pyarrow as pa
    return pa.Table.from_pylist(rows, schema=None) if rows else pa.table({c: pa.array([], pa.string()) for c in cols})


def table_bytes(fmt: str, cols, rows) -> bytes | None:
    """The table written in `fmt`, or None when no writer exists for it here."""
    if not cols:
        return b""
    if fmt == "csv":
        return _delimited(cols, rows, ",")
    if fmt == "tsv":
        return _delimited(cols, rows, "\t")
    if fmt in ("parquet", "orc", "feather", "arrow"):
        import pyarrow as pa
        t = _arrow_table(cols, rows)
        sink = io.BytesIO()
        if fmt == "parquet":
            import pyarrow.parquet as pq
            pq.write_table(t, sink)
        elif fmt == "orc":
            import pyarrow.orc as orc
            orc.write_table(t, sink)
        elif fmt == "feather":
            import pyarrow.feather as ft
            ft.write_feather(t, sink)
        else:
            with pa.ipc.new_file(sink, t.schema) as w:
                w.write_table(t)
        return sink.getvalue()
    if fmt == "avro":
        import fastavro
        def typ(c):
            vals = [r.get(c) for r in rows if r.get(c) is not None]
            base = ("boolean" if vals and all(isinstance(v, bool) for v in vals) else
                    "long" if vals and all(isinstance(v, int) and not isinstance(v, bool) for v in vals) else
                    "double" if vals and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals) else "string")
            return ["null", base]
        schema = fastavro.parse_schema({"type": "record", "name": "Row",
                                        "fields": [{"name": c, "type": typ(c)} for c in cols]})
        sink = io.BytesIO()
        fastavro.writer(sink, schema, rows)
        return sink.getvalue()
    if fmt in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(list(cols))
        for r in rows:
            ws.append([r.get(c) for c in cols])
        sink = io.BytesIO()
        wb.save(sink)
        return sink.getvalue()
    return None                                        # xls: no writer in this image


def avro_safe_names(cols) -> bool:
    import re
    return all(re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", c) for c in cols)


# ── documents ────────────────────────────────────────────────────────────────────────────────────
MARKER = "OpenCortyx regression marker sentence"
DOC_FORMATS = ("txt", "md", "markdown", "rst", "log", "hl7", "csl", "tex", "html", "htm", "eml",
               "msg", "pdf", "docx", "doc", "pptx", "odt")
DOC_VARIANTS = {
    "normal": ("accept", f"{MARKER}. Quarterly revenue grew in Oslo and Bergen."),
    "unicode": ("accept", f"{MARKER}. 東京の売上 — Zürich café — مرحبا"),
    "long_200k": ("accept", (f"{MARKER}. " + "Paragraph of steady filler text. " * 7000)),
    "empty": ("reject", ""),
    "whitespace": ("reject", "   \n\n\t  "),
}


def doc_bytes(fmt: str, text: str) -> bytes | None:
    if fmt in ("txt", "log", "csl"):
        return text.encode()
    if fmt in ("md", "markdown"):
        return (f"# Report\n\n{text}\n" if text.strip() else text).encode()
    if fmt == "rst":
        return (f"Report\n======\n\n{text}\n" if text.strip() else text).encode()
    if fmt == "tex":
        return (f"\\documentclass{{article}}\\begin{{document}}{text}\\end{{document}}" if text.strip() else text).encode()
    if fmt == "hl7":
        return (f"MSH|^~\\&|APP|FAC|||20240101||ORU^R01|1|P|2.5\rOBX|1|TX|NOTE||{text}" if text.strip() else text).encode()
    if fmt in ("html", "htm"):
        return (f"<html><head><title>Report</title></head><body><p>{text}</p></body></html>" if text.strip() else text).encode()
    if fmt == "eml":
        from email.message import EmailMessage
        m = EmailMessage()
        m["From"], m["To"], m["Subject"] = "a@example.test", "b@example.test", "Report"
        m.set_content(text)
        return bytes(m) if text.strip() else text.encode()
    if fmt == "pdf":
        import fitz
        d = fitz.open()
        if text.strip():
            page = d.new_page()
            page.insert_textbox(fitz.Rect(40, 40, 560, 800), text[:3000], fontsize=10)
        else:
            d.new_page()
        return d.tobytes()
    if fmt == "docx":
        import docx
        d = docx.Document()
        if text.strip():
            d.add_paragraph(text)
        s = io.BytesIO()
        d.save(s)
        return s.getvalue()
    if fmt == "pptx":
        import pptx
        p = pptx.Presentation()
        sl = p.slides.add_slide(p.slide_layouts[5])
        if text.strip():
            sl.shapes.title.text = text[:5000]
        s = io.BytesIO()
        p.save(s)
        return s.getvalue()
    if fmt == "odt":
        from xml.sax.saxutils import escape
        s = io.BytesIO()
        with zipfile.ZipFile(s, "w") as z:
            z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
            z.writestr("content.xml", '<?xml version="1.0" encoding="UTF-8"?><office:document-content '
                       'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                       'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body>'
                       f'<office:text><text:p>{escape(text)}</text:p></office:text></office:body>'
                       '</office:document-content>')
        return s.getvalue()
    return None                                        # doc, msg: legacy binary, no writer here


# ── chunking texts ───────────────────────────────────────────────────────────────────────────────
CHUNK_TEXTS = {
    "prose_short": "One sentence only.",
    "prose_paragraphs": "\n\n".join(f"Paragraph {i}. It has two sentences here." for i in range(40)),
    "sentences_varied": " ".join(("Short. " if i % 3 else "This sentence is considerably longer than the others around it. ") for i in range(120)),
    "markdown_sections": "Intro text.\n\n" + "\n\n".join(f"## Section {i}\nBody of section {i}. " + "More. " * (i * 4) for i in range(12)),
    "markdown_nested": "# Top\nA.\n## Mid\nB.\n### Low\nC.\n#### Lower\nD.\n",
    "no_headings_markdown": "Plain text without any heading line at all, only prose. " * 30,
    "single_long_word": "x" * 5000,
    "unicode_text": "東京は大きい。 " * 200 + "مرحبا بالعالم. " * 50,
    "whitespace_heavy": "\n\n\n   word   \n\n\n" * 50,
    "crlf_lines": "Line one.\r\nLine two.\r\n\r\nPara two.\r\n" * 40,
    "numbers_abbrev": "Dr. Smith paid $3.50 on 1.2.2024. Mr. Jones did not. " * 30,
}
CHUNK_STRATEGIES = ("sliding_window", "sentence", "paragraph", "markdown_heading", "unknown_strategy")
CHUNK_SIZES = (40, 200, 1000, 4000)
CHUNK_OVERLAPS = (0, 10, 39)

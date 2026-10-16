"""Tests for layout-aware document AST powered by Docling:
dynamic node types, 2D table grids, asset references, and dual-edge topology (CONTAINS + NEXT).
"""
import fitz

from opencortyx.ingest.formats.unstructured.docling import engine as docling_engine
from opencortyx.ingest.formats.unstructured.docling.assembler import DynamicDoclingAssembler
from opencortyx.ingest.formats.unstructured.docling.engine import DoclingEngine
from opencortyx.ingest.formats.unstructured.docling.schemas import to_pascal_case
from opencortyx.ingest.walkers.unstructured import UnstructuredWalker


def test_to_pascal_case_dynamic_labels():
    """Verify any Docling label translates dynamically into PascalCase node type."""
    assert to_pascal_case("section_header") == "SectionHeader"
    assert to_pascal_case("table") == "Table"
    assert to_pascal_case("paragraph") == "Paragraph"
    assert to_pascal_case("picture") == "Picture"
    assert to_pascal_case("footnote") == "Footnote"
    assert to_pascal_case("formula") == "Formula"
    assert to_pascal_case("code") == "Code"
    assert to_pascal_case("custom_complex_label_123") == "CustomComplexLabel123"


def test_dynamic_docling_assembler_hierarchy_and_next_flow():
    """DynamicDoclingAssembler tracks heading stack for parent_id and chains next_id sequentially."""
    assembler = DynamicDoclingAssembler(document_id="doc-42")

    # 1. Root heading H1
    h1 = assembler.add_element(
        label="section_header",
        text="Introduction",
        level=1,
        page_no=1,
    )
    assert h1.parent_id == "doc-42"
    assert h1.node_type == "SectionHeader"
    assert h1.next_id is None

    # 2. Paragraph under H1
    p1 = assembler.add_element(
        label="paragraph",
        text="This is the opening introduction paragraph.",
        page_no=1,
    )
    assert p1.parent_id == h1.id
    assert p1.node_type == "Paragraph"
    # h1.next_id must now point to p1!
    assert h1.next_id == p1.id

    # 3. Sub-heading H2
    h2 = assembler.add_element(
        label="section_header",
        text="System Architecture",
        level=2,
        page_no=1,
    )
    assert h2.parent_id == h1.id
    assert h2.node_type == "SectionHeader"
    # p1.next_id must point to h2!
    assert p1.next_id == h2.id

    # 4. Table under H2 with 2D grid matrix
    table_grid = [
        ["Header A", "Header B"],
        ["Val 1", "Val 2"],
    ]
    t1 = assembler.add_element(
        label="table",
        text="| Header A | Header B |\n| Val 1 | Val 2 |",
        grid=table_grid,
        page_no=2,
    )
    assert t1.parent_id == h2.id
    assert t1.node_type == "Table"
    assert t1.grid == table_grid
    # h2.next_id points to t1
    assert h2.next_id == t1.id

    # 5. Picture under H2 with asset reference
    pic = assembler.add_element(
        label="picture",
        text="Architecture diagram",
        asset_reference="data:image/png;base64,iVBORw0KGgo...",
        page_no=2,
    )
    assert pic.parent_id == h2.id
    assert pic.node_type == "Picture"
    assert pic.asset_reference == "data:image/png;base64,iVBORw0KGgo..."
    # t1.next_id points to pic
    assert t1.next_id == pic.id

    # Build and verify sequential ordering & unique types
    nodes = assembler.build()
    assert len(nodes) == 5
    assert assembler.get_unique_node_types() == {"SectionHeader", "Paragraph", "Table", "Picture"}
    # Verify linear chain: h1 -> p1 -> h2 -> t1 -> pic
    chain = [n.id for n in nodes]
    assert nodes[0].next_id == chain[1]
    assert nodes[1].next_id == chain[2]
    assert nodes[2].next_id == chain[3]
    assert nodes[3].next_id == chain[4]
    assert nodes[4].next_id is None


def test_unstructured_walker_layout_introspection():
    """UnstructuredWalker introspects layout sample and emits dynamic node types with CONTAINS and NEXT edges."""
    sample = (
        "# Overview\n\n"
        "Here is a paragraph describing OpenCortyx.\n\n"
        # A Markdown table needs its delimiter row; Docling now reads pipes without one as text.
        "| Service | Port |\n| --- | --- |\n| API | 8000 |\n| Web | 3000 |\n\n"
        "## Submodule\n"
        "- Item 1\n"
        "- Item 2\n"
    )
    walker = UnstructuredWalker()
    result = walker.introspect(sample, strategy="layout")

    types = result["types"]
    # Dynamic types discovered from text
    assert "SectionHeader" in types
    # "Text", Docling's own label — the fallback no longer invents "Paragraph", so the graph's
    # node types do not depend on which engine read the document.
    assert "Text" in types
    assert "Paragraph" not in types
    assert "Table" in types

    # Edge rules should include both CONTAINS and NEXT
    edge_rules = result["edge_rules"]
    contains_edges = [e for e in edge_rules if e["type"] == "CONTAINS"]
    next_edges = [e for e in edge_rules if e["type"] == "NEXT"]

    assert len(contains_edges) > 0
    assert len(next_edges) > 0

    # NEXT edge connects dynamic nodes via next_id
    sample_next = next_edges[0]
    assert sample_next["to"]["ref"] == "next_id"
    assert sample_next["to"]["kind"] == "context"


# ── regressions: each of these shipped broken once ─────────────────────────────────────────────
def _fallback_only(monkeypatch):
    """Force the structural fallback, as in any process where Docling cannot load."""
    monkeypatch.setattr(DoclingEngine, "_converter_instance", None)
    monkeypatch.setattr(DoclingEngine, "unavailable_reason", "forced by test")


def test_the_title_contains_the_sections_instead_of_standing_beside_them(monkeypatch):
    """Title and top-level sections were both level 1, so every section became the title's
    sibling and the title contained only the text before the first section."""
    _fallback_only(monkeypatch)
    a = DoclingEngine.convert(b"# Handbook\n\nIntro.\n\n## Leave\n\n25 days.\n\n### Sick\n\nTen.\n",
                              filename="handbook.md", doc_id="doc")
    by_text = {n.text: n for n in a.nodes}
    title, leave, sick = by_text["Handbook"], by_text["Leave"], by_text["Sick"]
    assert title.type == "Title" and title.parent_id == "doc"
    assert leave.parent_id == title.id, "a level-1 section belongs under the title"
    assert sick.parent_id == leave.id
    assert by_text["Ten."].parent_id == sick.id


def test_a_pdf_the_fallback_reads_never_becomes_pdf_syntax(monkeypatch):
    """Decoding PDF bytes as text produced 318 nodes reading "xref", "endobj", "%%EOF"."""
    _fallback_only(monkeypatch)
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Payment terms are thirty days.", fontsize=12)
    a = DoclingEngine.convert(doc.tobytes(), filename="terms.pdf", doc_id="doc")
    texts = " ".join(n.text for n in a.nodes)
    assert "thirty days" in texts
    for syntax in ("endobj", "xref", "%%EOF", "/Type"):
        assert syntax not in texts, f"PDF syntax {syntax!r} leaked into the graph"
    assert a.engine == "fallback" and a.fallback_reason


def test_a_converter_that_crashes_on_load_falls_back_instead_of_failing_the_ingest(monkeypatch):
    """docling under pydantic < 2.10 raises AttributeError, not ImportError, at import. Only
    ImportError was caught, so the error escaped and the whole document failed."""
    monkeypatch.setattr(DoclingEngine, "_converter_instance", None)
    monkeypatch.setattr(DoclingEngine, "unavailable_reason", None)

    def _boom(*_a, **_kw):
        raise AttributeError("force_full_page_ocr")
    monkeypatch.setattr("builtins.__import__", _import_raising("docling", _boom))
    a = DoclingEngine.convert(b"# T\n\nbody\n", filename="t.md", doc_id="doc")
    assert a.engine == "fallback"
    assert "AttributeError" in DoclingEngine.unavailable_reason
    assert [n.type for n in a.nodes] == ["Title", "Text"]


def test_a_table_without_text_is_given_its_grid_as_text():
    """Docling leaves TableItem.text empty; an empty-text node is never embedded or retrieved."""
    assert docling_engine._grid_text([["Item", "Limit"], ["Hotel", "150"]]) == \
        "Item | Limit\nHotel | 150"


def test_a_sample_with_no_extension_is_not_sent_to_docling_as_a_pdf():
    """The Mapper previews a text sample with no filename; `.pdf` was assumed, so it always failed."""
    assert docling_engine._effective_filename("# hi", "").endswith(".md")
    assert docling_engine._effective_filename(b"%PDF-1.7 ...", "").endswith(".pdf")
    assert docling_engine._effective_filename(b"x", "keep.docx") == "keep.docx"


def test_only_headings_open_a_section():
    """An element carrying a level (a list item, a caption) must not adopt what follows it."""
    a = DynamicDoclingAssembler(doc_id="doc")
    h = a.add_element(raw_label="section_header", text="Scope", level=1)
    li = a.add_element(raw_label="list_item", text="one", level=3)
    after = a.add_element(raw_label="text", text="after the list")
    assert li.parent_id == h.id and li.level is None
    assert after.parent_id == h.id, "the list item must not have become a heading"


def _import_raising(prefix, fn):
    real = __import__

    def _imp(name, *args, **kwargs):
        if name == prefix or name.startswith(prefix + "."):
            return fn()
        return real(name, *args, **kwargs)
    return _imp



# ── Page furniture and heading levels ────────────────────────────────────────────────────────────

class _Label:
    def __init__(self, value):
        self.value = value


class _Item:
    def __init__(self, label, text="", level=None):
        self.label, self.text, self.level = _Label(label), text, level


def test_pdf_heading_levels_come_from_section_numbers():
    """A PDF's layout model calls every heading level 1; the section number says how deep it is."""
    title, one, one_one, two, plain = (_Item("section_header", "Employee Handbook", 1),
                                      _Item("section_header", "1. Onboarding", 1),
                                      _Item("section_header", "1.1 First week", 1),
                                      _Item("section_header", "2 Benefits", 1),
                                      _Item("section_header", "Appendix", 1))
    levels = docling_engine._heading_levels([title, _Item("text", "intro"), one, one_one, two, plain])
    assert levels[id(title)] == 0                       # opens a numbered document: the title
    assert (levels[id(one)], levels[id(one_one)], levels[id(two)]) == (1, 2, 1)
    assert id(plain) not in levels                      # unnumbered later heading: as Docling gave it


def test_a_recorded_heading_level_or_title_is_kept():
    real = _Item("section_header", "1.1.1 Deep", 2)     # Word gave a level: it wins
    levels = docling_engine._heading_levels([_Item("title", "Handbook"), _Item("section_header", "Intro", 1), real])
    assert id(real) not in levels and not any(v == 0 for v in levels.values())


def _handbook_pdf() -> bytes:
    doc = fitz.open()
    for n, blocks in enumerate([[("Employee Handbook", 22), ("This handbook explains onboarding.", 10.5),
                                 ("1. Onboarding", 15), ("New hires complete IT setup on day one.", 10.5),
                                 ("1.1 First week", 12.5), ("Collect your laptop from the front desk.", 10.5)],
                                [("2. Leave", 15), ("Employees accrue 1.5 days of leave per month.", 10.5)]], 1):
        page = doc.new_page(width=595, height=842)
        page.insert_text((50, 35), "ACME Corp - Employee Handbook", fontsize=8)
        page.insert_text((50, 812), f"Confidential - Page {n}", fontsize=8)
        y = 90
        for text, size in blocks:
            page.insert_text((50, y), text, fontsize=size, fontname="hebo" if size > 11 else "helv")
            y += size * 2.2
    return doc.tobytes()


def test_a_pdf_keeps_its_page_headers_footers_and_section_tree():
    """Running headers and footers are read (they live in Docling's furniture layer) and belong to
    the document; numbered sections nest under the title and under each other."""
    a = DoclingEngine.convert(_handbook_pdf(), filename="handbook.pdf")
    if a.engine != "docling":
        import pytest
        pytest.skip(f"Docling unavailable here: {a.fallback_reason}")
    rows = {r["text"]: r for r in a.to_dict_list()}
    assert [r["label"] for r in a.to_dict_list()].count("page_header") == 2
    assert [r["label"] for r in a.to_dict_list()].count("page_footer") == 2
    for r in a.to_dict_list():
        if r["label"] in ("page_header", "page_footer"):
            assert r["parent_id"] == a.doc_id
    title = rows["Employee Handbook"]
    assert title["label"] == "title" and title["level"] == 0
    assert rows["1. Onboarding"]["parent_id"] == title["id"]
    assert rows["1.1 First week"]["parent_id"] == rows["1. Onboarding"]["id"]
    assert rows["Collect your laptop from the front desk."]["parent_id"] == rows["1.1 First week"]["id"]


def test_a_word_documents_header_and_footer_are_read():
    import io

    import docx

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "ACME Corp - Employee Handbook"
    d.sections[0].footer.paragraphs[0].text = "Confidential"
    d.add_heading("Employee Handbook", level=0)
    d.add_heading("1. Onboarding", level=1)
    d.add_paragraph("New hires complete IT setup on day one.")
    buf = io.BytesIO()
    d.save(buf)
    a = DoclingEngine.convert(buf.getvalue(), filename="handbook.docx")
    if a.engine != "docling":
        import pytest
        pytest.skip(f"Docling unavailable here: {a.fallback_reason}")
    labels = {r["text"]: r["label"] for r in a.to_dict_list()}
    assert labels["ACME Corp - Employee Handbook"] == "page_header"
    assert labels["Confidential"] == "page_footer"
